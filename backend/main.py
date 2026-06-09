import os
import logging
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from typing import List, Optional
from pydantic import BaseModel

# Import the agent from our agent module
import sys
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from agent.agent import agent
from google.adk import Runner
from google.adk.sessions.in_memory_session_service import InMemorySessionService
from google.genai import types

# Setup logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("reporting-backend")

app = FastAPI(title="Reporting Agent API")

from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse

# ... (after app initialization and CORS setup)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"], 
    allow_credentials=True,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["*"],
)

# Initialize ADK components
session_service = InMemorySessionService()
runner = Runner(
    agent=agent, 
    session_service=session_service, 
    app_name="reporting_app", 
    auto_create_session=True
)

# Serve the React frontend static files
# Make sure the 'dist' folder exists after 'npm run build'
frontend_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "frontend", "dist")

if os.path.exists(frontend_path):
    app.mount("/assets", StaticFiles(directory=os.path.join(frontend_path, "assets")), name="assets")
    
    @app.get("/")
    async def serve_frontend():
        return FileResponse(os.path.join(frontend_path, "index.html"))
else:
    @app.get("/")
    async def root():
        return {"status": "online", "message": "Backend online, but frontend not built. Run 'npm run build' in the frontend folder."}

import io
import base64
from pptx import Presentation
from pypdf import PdfReader

@app.post("/report")
async def generate_report(
    message: str = Form(...),
    user_id: str = Form("default_user"),
    session_id: str = Form("default_session"),
    files: Optional[List[UploadFile]] = File(None)
):
    try:
        # 1. Process uploaded files
        text_parts = [message]
        multimodal_parts = []
        pdf_base64 = None
        
        if files:
            for file in files:
                filename = file.filename or "unknown"
                mime_type = file.content_type or ""
                content = await file.read()
                
                # Capture the first PDF for preview
                if (filename.endswith(".pdf") or mime_type == "application/pdf") and pdf_base64 is None:
                    pdf_base64 = base64.b64encode(content).decode('utf-8')

                if filename.endswith(".pptx"):
                    try:
                        prs = Presentation(io.BytesIO(content))
                        pptx_text = [f"--- Content from PPTX '{filename}' ---"]
                        for i, slide in enumerate(prs.slides):
                            pptx_text.append(f"Slide {i+1}:")
                            for shape in slide.shapes:
                                if hasattr(shape, "text"):
                                    pptx_text.append(shape.text)
                        text_parts.append("\n".join(pptx_text))
                    except Exception as e:
                        text_parts.append(f"\n[Error extracting PPTX '{filename}': {str(e)}]")
                
                elif filename.endswith(".pdf") or mime_type == "application/pdf":
                    try:
                        reader = PdfReader(io.BytesIO(content))
                        pdf_text = [f"--- Content from PDF '{filename}' ---"]
                        for i, page in enumerate(reader.pages):
                            pdf_text.append(f"Page {i+1}:\n" + (page.extract_text() or ""))
                        text_parts.append("\n".join(pdf_text))
                    except Exception as e:
                        text_parts.append(f"\n[Error extracting PDF '{filename}': {str(e)}]")
                
                elif mime_type.startswith("image/"):
                    # Add as a multimodal part for Gemini Vision
                    multimodal_parts.append(
                        types.Part(
                            inline_data=types.Blob(
                                mime_type=mime_type,
                                data=content
                            )
                        )
                    )
                    text_parts.append(f"\n[Image file '{filename}' attached for visual analysis]")
                
                else:
                    # Fallback for text files
                    try:
                        text_content = content.decode("utf-8")
                        text_parts.append(f"\nContent from uploaded file '{filename}':\n{text_content}")
                    except:
                        text_parts.append(f"\n[Uploaded file '{filename}' (unsupported binary format)]")

        # Combine text context
        full_text_query = "\n".join(text_parts)
        
        # Build the final message parts
        final_parts = [types.Part(text=full_text_query)] + multimodal_parts
        
        # 2. Run the agent
        events = runner.run(
            user_id=user_id,
            session_id=session_id,
            new_message=types.Content(parts=final_parts)
        )
        
        # 3. Collect response
        full_response = ""
        for event in events:
            if event.content and event.content.parts:
                for part in event.content.parts:
                    if part.text:
                        full_response += part.text
        
        return {
            "response": full_response,
            "session_id": session_id,
            "raw_context": full_text_query,
            "pdf_data": pdf_base64
        }
        
    except Exception as e:
        logger.error(f"Error in /report: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8001)
